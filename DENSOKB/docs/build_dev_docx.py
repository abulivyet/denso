# -*- coding: utf-8 -*-
"""Generate a formal Word development specification from the markdown source."""
from __future__ import annotations

import re
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from formal_docx import (
    add_body,
    add_code_block,
    add_heading,
    add_list_item,
    add_table,
    add_toc,
    add_unnumbered_title,
    build_cover,
    center_cols_for,
    clean_md,
    new_document,
    setup_cover_section,
    start_body_section,
    table_widths,
)


def is_table_sep(line):
    s = line.strip().strip("|")
    if not s:
        return False
    parts = [p.strip() for p in s.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", p) for p in parts if p)


def parse_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [clean_md(c.strip()) for c in s.split("|")]


def normalize_heading(text):
    text = clean_md(text)
    text = re.sub(r"^(\d+(?:\.\d+)*)\.\s+", r"\1  ", text)
    return text


def convert_markdown(doc, lines, caption_start=1):
    caption_no = caption_start
    last_heading = ""
    caption_used = {}
    i = 0
    in_code = False
    code_buf = []
    ol_n = 0

    while i < len(lines):
        line = lines[i]
        raw = line.rstrip("\n")

        if raw.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_buf)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(raw)
            i += 1
            continue

        if not raw.strip() or raw.strip() == "---":
            ol_n = 0
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", raw)
        if m:
            level = len(m.group(1))
            text = normalize_heading(m.group(2))
            if level == 1:
                i += 1
                continue
            add_heading(doc, text, min(level - 1, 3))
            last_heading = re.sub(r"^\d+(?:\.\d+)*\s+", "", text)
            ol_n = 0
            i += 1
            continue

        if raw.lstrip().startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [parse_table_row(raw)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            headers, data = rows[0], rows[1:]
            cap = last_heading or "数据表"
            caption_used[cap] = caption_used.get(cap, 0) + 1
            if caption_used[cap] > 1:
                cap = f"{cap}（{caption_used[cap]}）"
            add_table(
                doc,
                headers,
                data,
                table_widths(headers),
                center_cols=center_cols_for(headers),
                caption=cap,
                caption_no=caption_no,
            )
            caption_no += 1
            ol_n = 0
            continue

        ul = re.match(r"^(\s*)[-*]\s+(.*)$", raw)
        if ul:
            ol_n = 0
            add_list_item(doc, ul.group(2).rstrip(), "· ")
            i += 1
            continue

        ol = re.match(r"^(\s*)(\d+)\.\s+(.*)$", raw)
        if ol:
            ol_n = int(ol.group(2)) if ol_n == 0 else ol_n + 1
            add_list_item(doc, ol.group(3).rstrip(), f"{ol.group(2)}. ")
            i += 1
            continue

        ol_n = 0
        text = raw.strip()
        if text.startswith("**") and text.endswith("**") and text.count("**") == 2:
            add_body(doc, text, first_indent=False)
        else:
            add_body(doc, text, first_indent=True)
        i += 1

    return caption_no


def build_from_md(
    md_path: Path,
    out_path: Path,
    *,
    title,
    english_title,
    header_right,
    doc_no,
    revision_note,
    cover_meta_keys,
    comments="",
):
    lines = md_path.read_text(encoding="utf-8").splitlines()

    meta = []
    body_start = 0
    for i, line in enumerate(lines):
        m = re.match(r"^\*\*([^：:*]+)[：:]\*\*\s*(.*)$", line.strip())
        if m:
            meta.append((m.group(1).strip(), clean_md(m.group(2))))
            body_start = i + 1
        elif line.strip() == "---" and meta:
            body_start = i + 1
            break
        elif line.startswith("## "):
            body_start = i
            break

    doc = new_document()
    setup_cover_section(doc.sections[0])
    doc.core_properties.title = "南沙电装 OMS 系统" + title
    doc.core_properties.subject = "DENSOKB"
    doc.core_properties.author = "南沙电装 OMS"
    doc.core_properties.comments = comments or (title + " V1.0")

    cover_meta = [(k, v) for k, v in meta if k in cover_meta_keys]
    build_cover(
        doc,
        chinese_title=title,
        english_title=english_title,
        meta_rows=cover_meta,
        doc_no=doc_no,
        version="V1.0",
        date="2026年9月3日",
    )

    start_body_section(doc, header_right)

    add_unnumbered_title(doc, "修订记录")
    add_body(doc, "本文档随系统功能变更同步修订。历次修订记录如下。")
    add_table(
        doc,
        ["版本", "日期", "修订说明", "编制"],
        [["V1.0", "2026-09-03", revision_note, ""]],
        [2.0, 2.6, 8.4, 2.5],
        center_cols={0, 1, 3},
        caption="修订记录",
        caption_no=1,
    )
    doc.add_page_break()
    add_toc(doc)

    convert_markdown(doc, lines[body_start:], caption_start=2)

    from formal_docx import add_text, set_paragraph_format
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Cm

    p = doc.add_paragraph()
    set_paragraph_format(p, before=24, after=0, line=28, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "— 全文完 —", east="宋体", size=12)

    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(out_path))
    print("wrote", out_path)


def build(md_path: Path, out_path: Path):
    build_from_md(
        md_path,
        out_path,
        title="系统开发说明书",
        english_title="Software Development Specification",
        header_right="系统开发说明书",
        doc_no="OMS-DEV-SPEC-2026-001",
        revision_note="初稿。依据现行代码、配置与页面整理系统结构、流程、接口及开发约定。",
        cover_meta_keys=("系统名称", "工程名", "入口页"),
        comments="系统开发说明书 V1.0",
    )


if __name__ == "__main__":
    docs = Path(__file__).resolve().parent
    build(docs / "南沙电装OMS-开发文档.md", docs / "南沙电装OMS-开发文档.docx")
    build_from_md(
        docs / "南沙电装OMS-操作手册.md",
        docs / "南沙电装OMS-操作手册.docx",
        title="操作手册",
        english_title="User Operation Manual",
        header_right="操作手册",
        doc_no="OMS-OPS-MAN-2026-001",
        revision_note="初稿。依据现行页面功能整理登录、上传、各客户下载打印及常见问题。",
        cover_meta_keys=("系统名称", "适用对象"),
        comments="操作手册 V1.0",
    )
