# -*- coding: utf-8 -*-
"""Formal landscape Word for the operations issue log."""
from __future__ import annotations

import sys
from pathlib import Path

from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt

sys.path.insert(0, str(Path(__file__).resolve().parent))

import formal_docx as f
from build_dev_docx import build_from_md, convert_markdown


def set_landscape(section):
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Cm(29.7)
    section.page_height = Cm(21.0)
    section.top_margin = Cm(1.8)
    section.bottom_margin = Cm(1.8)
    section.left_margin = Cm(2.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.2)
    section.footer_distance = Cm(1.2)


def landscape_widths(headers):
    n = len(headers)
    if n == 2:
        return [5.0, 20.7]
    if n == 5:
        return [1.4, 2.2, 6.4, 7.6, 8.1]
    w = round(25.7 / n, 2)
    widths = [w] * n
    widths[-1] = round(25.7 - w * (n - 1), 2)
    return widths


def build():
    docs = Path(__file__).resolve().parent
    f.CONTENT_CM = 25.7
    f.CONTENT_WIDTH = Cm(25.7)
    orig_widths = f.table_widths
    f.table_widths = landscape_widths

    out = docs / "南沙电装OMS-运维问题记录.docx"
    build_from_md(
        docs / "南沙电装OMS-运维问题记录.md",
        out,
        title="运维问题记录",
        english_title="Operation Issue Log",
        header_right="运维问题记录",
        doc_no="OMS-OPS-ISSUE-2026-001",
        revision_note="V1.0 初稿整理当年现场问题；V1.1 补记 2026 年 3 月已处理未建档问题（序号 11～24）。",
        cover_meta_keys=("整理说明", "客户简称"),
        comments="运维问题记录 V1.1",
    )

    from docx import Document

    doc = Document(str(out))
    for section in doc.sections:
        set_landscape(section)
    doc.save(str(out))
    f.table_widths = orig_widths
    print("landscape", out)


if __name__ == "__main__":
    # build_from_md already writes portrait; rebuild body with landscape widths instead
    docs = Path(__file__).resolve().parent
    f.CONTENT_CM = 25.7
    f.CONTENT_WIDTH = Cm(25.7)
    f.table_widths = landscape_widths

    md = docs / "南沙电装OMS-运维问题记录.md"
    out = docs / "南沙电装OMS-运维问题记录.docx"
    lines = md.read_text(encoding="utf-8").splitlines()

    doc = f.new_document()
    set_landscape(doc.sections[0])
    f.setup_cover_section(doc.sections[0])
    doc.core_properties.title = "南沙电装 OMS 系统运维问题记录"
    doc.core_properties.author = "南沙电装 OMS"

    f.build_cover(
        doc,
        chinese_title="运维问题记录",
        english_title="Operation Issue Log",
        meta_rows=[("客户范围", "GHAC / DHEC / CHAM / FTEC"), ("版本", "V1.1（补 3 月问题）")],
        doc_no="OMS-OPS-ISSUE-2026-001",
        version="V1.1",
        date="2026年9月18日",
    )

    body = f.start_body_section(doc, "运维问题记录")
    set_landscape(body)

    f.add_unnumbered_title(doc, "修订记录")
    f.add_body(doc, "本文档记录现场已处理问题及程序侧待改点，便于交接和后续改造。")
    f.add_table(
        doc,
        ["版本", "日期", "修订说明", "编制"],
        [
            ["V1.0", "2026-09-03", "初稿。对照代码整理当年现场问题（序号 1～10）。", ""],
            ["V1.1", "2026-09-18", "补记 2026 年 3 月已处理但未建档的问题（序号 11～24）。", ""],
        ],
        [2.4, 3.0, 16.3, 4.0],
        center_cols={0, 1, 3},
        caption="修订记录",
        caption_no=1,
    )
    doc.add_page_break()
    f.add_toc(doc)

    body_start = 0
    for i, line in enumerate(lines):
        if line.startswith("## "):
            body_start = i
            break
    convert_markdown(doc, lines[body_start:], caption_start=2)

    p = doc.add_paragraph()
    f.set_paragraph_format(
        p, before=24, after=0, line=28, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER
    )
    f.add_text(p, "— 全文完 —", east="宋体", size=12)

    doc.save(str(out))
    print("wrote", out)
