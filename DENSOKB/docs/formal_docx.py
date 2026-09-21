# -*- coding: utf-8 -*-
"""Shared formal Word styles for OMS manuals (A4, 小四, 固定 28 磅)."""
from __future__ import annotations

import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

CONTENT_CM = 15.5
CONTENT_WIDTH = Cm(CONTENT_CM)
NAVY = "1F4E79"
NAVY_RGB = RGBColor(0x1F, 0x4E, 0x79)
HEADER_FILL = "1F4E79"
ALT_FILL = "F5F7FA"
CODE_FILL = "F4F6F8"
BLACK = RGBColor(0x00, 0x00, 0x00)
BODY = RGBColor(0x22, 0x22, 0x22)
CODE_COLOR = RGBColor(0x1F, 0x4E, 0x79)


def set_run_font(run, east="宋体", latin="Times New Roman", size=12, bold=False, color=None):
    run.bold = bold
    run.italic = False
    run.font.size = Pt(size)
    run.font.name = latin
    if color is not None:
        run.font.color.rgb = color
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), latin)
    rFonts.set(qn("w:hAnsi"), latin)
    rFonts.set(qn("w:eastAsia"), east)
    rFonts.set(qn("w:cs"), latin)


def add_text(paragraph, text, **kwargs):
    run = paragraph.add_run(text)
    set_run_font(run, **kwargs)
    return run


def set_paragraph_format(
    p,
    before=0,
    after=0,
    line=None,
    exact=False,
    first_indent=None,
    align=None,
    keep_next=False,
):
    pf = p.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    if line is not None:
        if exact:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(line)
        else:
            pf.line_spacing = line
    if first_indent is not None:
        pf.first_line_indent = first_indent
    if align is not None:
        p.alignment = align
    pf.keep_with_next = keep_next


def apply_page(section):
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.5)
    section.footer_distance = Cm(1.75)


def new_document():
    doc = Document()
    apply_page(doc.sections[0])
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.font.color.rgb = BODY
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    nf = normal.paragraph_format
    nf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    nf.line_spacing = Pt(28)
    nf.space_before = Pt(0)
    nf.space_after = Pt(0)
    el = OxmlElement("w:updateFields")
    el.set(qn("w:val"), "true")
    doc.settings.element.append(el)
    return doc


def shade_cell(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def set_cell_borders(cell, color="8FA2B5", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    old = tcPr.find(qn("w:tcBorders"))
    if old is not None:
        tcPr.remove(old)
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), sz)
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def set_row_height(row, cm, rule="atLeast"):
    trPr = row._tr.get_or_add_trPr()
    old = trPr.find(qn("w:trHeight"))
    if old is not None:
        trPr.remove(old)
    el = OxmlElement("w:trHeight")
    el.set(qn("w:val"), str(int(Cm(cm).twips)))
    el.set(qn("w:hRule"), rule)
    trPr.append(el)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader")
    el.set(qn("w:val"), "true")
    trPr.append(el)


def prevent_row_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def set_table_layout(table, widths_cm):
    tbl = table._tbl
    tblPr = tbl.tblPr
    tblW = tblPr.find(qn("w:tblW"))
    if tblW is None:
        tblW = OxmlElement("w:tblW")
        tblPr.append(tblW)
    total = int(sum(Cm(w).twips for w in widths_cm))
    tblW.set(qn("w:w"), str(total))
    tblW.set(qn("w:type"), "dxa")

    layout = tblPr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tblPr.append(layout)
    layout.set(qn("w:type"), "fixed")

    mar = tblPr.find(qn("w:tblCellMar"))
    if mar is not None:
        tblPr.remove(mar)
    mar = OxmlElement("w:tblCellMar")
    for edge, val in (("top", "40"), ("left", "80"), ("bottom", "40"), ("right", "80")):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:w"), val)
        el.set(qn("w:type"), "dxa")
        mar.append(el)
    tblPr.append(mar)

    grid = tbl.find(qn("w:tblGrid"))
    if grid is not None:
        tbl.remove(grid)
    grid = OxmlElement("w:tblGrid")
    for w in widths_cm:
        gc = OxmlElement("w:gridCol")
        gc.set(qn("w:w"), str(int(Cm(w).twips)))
        grid.append(gc)
    tbl.insert(1, grid)

    for row in table.rows:
        for i, cell in enumerate(row.cells):
            tcPr = cell._tc.get_or_add_tcPr()
            tcW = tcPr.find(qn("w:tcW"))
            if tcW is None:
                tcW = OxmlElement("w:tcW")
                tcPr.append(tcW)
            tcW.set(qn("w:w"), str(int(Cm(widths_cm[i]).twips)))
            tcW.set(qn("w:type"), "dxa")


def fill_cell(cell, text, size=10.5, bold=False, align="left", header=False, east="宋体"):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    p = cell.paragraphs[0]
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    pf = p.paragraph_format
    pf.space_before = Pt(2)
    pf.space_after = Pt(2)
    pf.line_spacing = 1.0
    pf.first_line_indent = Cm(0)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "right":
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    color = RGBColor(0xFF, 0xFF, 0xFF) if header else BODY
    add_text(p, text or "", east=east, size=size, bold=bold, color=color, latin="Times New Roman")


def add_table(doc, headers, rows, widths, center_cols=None, caption=None, caption_no=None):
    center_cols = set(center_cols or [])
    if caption:
        p = doc.add_paragraph()
        set_paragraph_format(
            p, before=6, after=3, line=16, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER
        )
        title = f"表 {caption_no}  {caption}" if caption_no else caption
        add_text(p, title, east="黑体", size=10.5, bold=True, color=BLACK)

    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_table_layout(table, widths)

    set_row_height(table.rows[0], 0.92)
    set_repeat_header(table.rows[0])
    prevent_row_split(table.rows[0])
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        shade_cell(cell, HEADER_FILL)
        set_cell_borders(cell, "1F4E79", "4")
        fill_cell(cell, h, size=10.5, bold=True, align="center", header=True, east="黑体")

    for r_idx, row_data in enumerate(rows):
        row = table.rows[r_idx + 1]
        set_row_height(row, 0.78)
        prevent_row_split(row)
        for c_idx in range(len(headers)):
            cell = row.cells[c_idx]
            val = row_data[c_idx] if c_idx < len(row_data) else ""
            if r_idx % 2 == 1:
                shade_cell(cell, ALT_FILL)
            set_cell_borders(cell, "B7C4D1", "4")
            align = "center" if c_idx in center_cols else "left"
            fill_cell(cell, str(val), size=10.5, bold=False, align=align)

    spacer = doc.add_paragraph()
    set_paragraph_format(spacer, before=0, after=6, line=6, exact=True, first_indent=Cm(0))
    return table


def add_mixed_runs(paragraph, text, size=12, first_code=False):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            add_text(paragraph, text[pos : m.start()], east="宋体", size=size, color=BODY)
        token = m.group(0)
        if token.startswith("**"):
            add_text(paragraph, token[2:-2], east="宋体", size=size, bold=True, color=BODY)
        else:
            add_text(
                paragraph,
                token[1:-1],
                east="宋体",
                latin="Consolas",
                size=size - 0.5 if size >= 11 else size,
                color=CODE_COLOR,
            )
        pos = m.end()
    if pos < len(text):
        add_text(paragraph, text[pos:], east="宋体", size=size, color=BODY)


def add_body(doc, text, first_indent=True):
    p = doc.add_paragraph()
    set_paragraph_format(
        p,
        before=0,
        after=0,
        line=28,
        exact=True,
        first_indent=Pt(24) if first_indent else Cm(0),
        align=WD_ALIGN_PARAGRAPH.JUSTIFY,
    )
    add_mixed_runs(p, text, size=12)
    return p


def add_unnumbered_title(doc, text):
    p = doc.add_paragraph()
    add_text(p, text, east="黑体", size=16, bold=True, color=BLACK)
    set_paragraph_format(
        p, before=18, after=12, line=28, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER
    )
    return p


def add_heading(doc, text, level):
    p = doc.add_heading("", level=level)
    for r in list(p.runs):
        r._element.getparent().remove(r._element)
    sizes = {1: 16, 2: 14, 3: 12, 4: 12}
    before = {1: 18, 2: 12, 3: 6, 4: 6}
    after = {1: 8, 2: 6, 3: 3, 4: 3}
    line = {1: 28, 2: 24, 3: 22, 4: 22}
    add_text(p, text, east="黑体", size=sizes[level], bold=True, color=BLACK)
    set_paragraph_format(
        p, before=before[level], after=after[level], line=line[level], exact=True, first_indent=Cm(0), keep_next=True
    )
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def add_list_item(doc, text, prefix):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
    pf.line_spacing = Pt(28)
    pf.left_indent = Cm(0.74)
    pf.first_line_indent = Cm(-0.74)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_text(p, prefix, east="宋体", size=12, color=BODY)
    add_mixed_runs(p, text, size=12)
    return p


def add_code_block(doc, lines):
    text = "\n".join(lines).rstrip()
    p = doc.add_paragraph()
    set_paragraph_format(p, before=6, after=6, line=14, exact=True, first_indent=Cm(0))
    p.paragraph_format.left_indent = Cm(0.3)
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), CODE_FILL)
    pPr.append(shd)
    add_text(p, text, east="宋体", latin="Consolas", size=9, color=CODE_COLOR)
    return p


def add_field_run(paragraph, instr, placeholder="1"):
    r1 = paragraph.add_run()
    b = OxmlElement("w:fldChar")
    b.set(qn("w:fldCharType"), "begin")
    r1._r.append(b)
    r2 = paragraph.add_run()
    it = OxmlElement("w:instrText")
    it.set(qn("xml:space"), "preserve")
    it.text = instr
    r2._r.append(it)
    r3 = paragraph.add_run()
    s = OxmlElement("w:fldChar")
    s.set(qn("w:fldCharType"), "separate")
    r3._r.append(s)
    r4 = paragraph.add_run(placeholder)
    set_run_font(r4, east="宋体", size=9, color=BODY)
    r5 = paragraph.add_run()
    e = OxmlElement("w:fldChar")
    e.set(qn("w:fldCharType"), "end")
    r5._r.append(e)


def _header_border(paragraph, edge="bottom", sz="12"):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement("w:" + edge)
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), sz)
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), NAVY)
    pBdr.append(el)
    pPr.append(pBdr)


def setup_cover_section(section):
    section.different_first_page_header_footer = False
    for part in (section.header, section.footer):
        part.is_linked_to_previous = False
        p = part.paragraphs[0]
        for r in list(p.runs):
            r._element.getparent().remove(r._element)
    fp = section.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_paragraph_format(fp, before=0, after=0, line=16, exact=True, first_indent=Cm(0))
    add_text(fp, "内部资料", east="宋体", size=9, color=BODY)


def setup_body_section(section, header_right):
    section.different_first_page_header_footer = False
    sectPr = section._sectPr
    pg = sectPr.find(qn("w:pgNumType"))
    if pg is None:
        pg = OxmlElement("w:pgNumType")
        sectPr.append(pg)
    pg.set(qn("w:start"), "1")

    header = section.header
    header.is_linked_to_previous = False
    hp = header.paragraphs[0]
    for r in list(hp.runs):
        r._element.getparent().remove(r._element)
    set_paragraph_format(hp, before=0, after=0, line=16, exact=True, first_indent=Cm(0))
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    add_text(hp, "南沙电装 OMS 系统", east="宋体", size=9, color=NAVY_RGB)
    hp.paragraph_format.tab_stops.add_tab_stop(CONTENT_WIDTH, WD_TAB_ALIGNMENT.RIGHT)
    add_text(hp, "\t" + header_right, east="宋体", size=9, color=NAVY_RGB)
    _header_border(hp, "bottom", "12")

    footer = section.footer
    footer.is_linked_to_previous = False
    fp = footer.paragraphs[0]
    for r in list(fp.runs):
        r._element.getparent().remove(r._element)
    set_paragraph_format(fp, before=0, after=0, line=16, exact=True, first_indent=Cm(0))
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _header_border(fp, "top", "6")
    add_text(fp, "第 ", east="宋体", size=9, color=BODY)
    add_field_run(fp, " PAGE ", "1")
    add_text(fp, " 页    共 ", east="宋体", size=9, color=BODY)
    add_field_run(fp, " SECTIONPAGES ", "1")
    add_text(fp, " 页", east="宋体", size=9, color=BODY)


def build_cover(doc, chinese_title, english_title, meta_rows, doc_no, version="V1.0", date="2026年9月3日"):
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=22, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.RIGHT)
    add_text(p, "密　　级：内部", east="黑体", size=12, bold=True, color=BLACK)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=48, after=0, line=28, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "南沙电装 OMS 系统", east="黑体", size=22, bold=True, color=NAVY_RGB)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=18, after=0, line=40, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, chinese_title, east="黑体", size=28, bold=True, color=BLACK)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=8, after=0, line=22, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, english_title, latin="Times New Roman", east="Times New Roman", size=14, color=NAVY_RGB)

    line_p = doc.add_paragraph()
    set_paragraph_format(line_p, before=12, after=18, line=6, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    pPr = line_p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "18")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), NAVY)
    pBdr.append(bottom)
    pPr.append(pBdr)

    for k, v in [("文档编号", doc_no), ("版　　本", version), ("编制日期", date)] + meta_rows:
        p = doc.add_paragraph()
        set_paragraph_format(p, before=0, after=0, line=28, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
        add_text(p, f"{k}：{v}", east="宋体", size=12, color=BODY)

    p = doc.add_paragraph()
    set_paragraph_format(p, before=36, after=6, line=22, exact=True, first_indent=Cm(0), align=WD_ALIGN_PARAGRAPH.CENTER)
    add_text(p, "签署栏", east="黑体", size=12, bold=True, color=BLACK)

    add_table(
        doc,
        ["项　　目", "姓　　名", "日　　期", "签　　名"],
        [
            ["编　　制", "", date, ""],
            ["审　　核", "", "", ""],
            ["批　　准", "", "", ""],
        ],
        [3.5, 4.0, 4.0, 4.0],
        center_cols={0, 1, 2, 3},
    )


def start_body_section(doc, header_right):
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    apply_page(section)
    setup_body_section(section, header_right)
    return section


def add_toc(doc):
    add_unnumbered_title(doc, "目　　录")
    p = doc.add_paragraph()
    set_paragraph_format(p, before=6, after=6, line=22, exact=True, first_indent=Cm(0))
    add_text(p, "打开文档后如页码未刷新，请全选后按 F9 更新域。", east="宋体", size=9, color=RGBColor(0x66, 0x66, 0x66))
    p = doc.add_paragraph()
    set_paragraph_format(p, before=0, after=0, line=28, exact=True, first_indent=Cm(0))
    add_field_run(p, r' TOC \o "1-3" \h \z \u ', "（目录域，打开 Word 后自动生成）")
    doc.add_page_break()


def table_widths(headers):
    n = len(headers)
    heads = [re.sub(r"[`*]", "", h).strip() for h in headers]
    if n == 2:
        return [4.2, 11.3]
    if n == 3:
        if heads[1] in ("路径", "脚本", "提交 type 片段", "type", "路径字段"):
            return [2.6, 6.2, 6.7]
        if heads[0] in ("条件", "代码", "Cust 参数", "配置键", "类", "键", "根", "界面", "fileType"):
            return [4.0, 4.8, 6.7]
        return [3.6, 5.4, 6.5]
    if n == 4:
        if "量产输入" in heads or "其它输入" in heads:
            return [1.8, 2.6, 4.4, 6.7]
        return [2.3, 3.6, 4.4, 5.2]
    if n == 5:
        return [1.8, 2.4, 3.5, 3.5, 4.3]
    w = round(CONTENT_CM / n, 2)
    widths = [w] * n
    widths[-1] = round(CONTENT_CM - w * (n - 1), 2)
    return widths


def center_cols_for(headers):
    cols = set()
    for i, h in enumerate(headers):
        t = re.sub(r"[`*]", "", h).strip()
        if t in ("序号", "方法", "代码", "项"):
            cols.add(i)
    return cols


def clean_md(text):
    text = re.sub(r"\*\*([^*]+)\*\*", r"\1", text)
    text = re.sub(r"`([^`]+)`", r"\1", text)
    return text.strip()
