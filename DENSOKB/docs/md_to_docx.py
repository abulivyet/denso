# -*- coding: utf-8 -*-
"""Convert project markdown docs to Word."""
import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


def set_run_font(run, name="宋体", size=11, bold=False, code=False):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    if code:
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.color.rgb = RGBColor(0xC4, 0x56, 0x26)
    else:
        run.font.name = name
        run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def add_formatted_runs(paragraph, text, size=11):
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    pos = 0
    for m in pattern.finditer(text):
        if m.start() > pos:
            run = paragraph.add_run(text[pos : m.start()])
            set_run_font(run, size=size)
        token = m.group(0)
        if token.startswith("**"):
            run = paragraph.add_run(token[2:-2])
            set_run_font(run, size=size, bold=True)
        else:
            run = paragraph.add_run(token[1:-1])
            set_run_font(run, size=size, code=True)
        pos = m.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run_font(run, size=size)


def is_table_sep(line):
    s = line.strip().strip("|")
    if not s:
        return False
    parts = [p.strip() for p in s.split("|")]
    return all(re.fullmatch(r":?-{3,}:?", p) for p in parts if p)


def parse_table_row(line):
    s = line.strip()
    if s.startswith("|"):
        s = s[1:]
    if s.endswith("|"):
        s = s[:-1]
    return [c.strip() for c in s.split("|")]


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement("w:" + edge)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), "BFBFBF")
        borders.append(el)
    table._tbl.tblPr.append(borders)


def convert(md_path: Path, docx_path: Path):
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()

    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.4)
    section.right_margin = Cm(2.4)

    styles = doc.styles
    styles["Normal"].font.name = "宋体"
    styles["Normal"].font.size = Pt(11)
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")

    heading_sizes = {1: 22, 2: 16, 3: 14, 4: 12}
    i = 0
    first_h1 = True
    in_code = False
    code_buf = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(8)
                p.paragraph_format.left_indent = Cm(0.3)
                run = p.add_run("\n".join(code_buf))
                set_run_font(run, size=9, code=True)
                code_buf = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if not line.strip():
            i += 1
            continue

        if line.strip() == "---":
            i += 1
            continue

        m = re.match(r"^(#{1,4})\s+(.*)$", line)
        if m:
            level = len(m.group(1))
            text = m.group(2).strip()
            p = doc.add_heading("", level=level)
            run = p.add_run(text)
            set_run_font(run, name="微软雅黑", size=heading_sizes[level], bold=True)
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            if first_h1 and level == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                first_h1 = False
            i += 1
            continue

        if line.lstrip().startswith("|") and i + 1 < len(lines) and is_table_sep(lines[i + 1]):
            rows = [parse_table_row(line)]
            i += 2
            while i < len(lines) and lines[i].lstrip().startswith("|"):
                rows.append(parse_table_row(lines[i]))
                i += 1
            cols = max(len(r) for r in rows)
            table = doc.add_table(rows=len(rows), cols=cols)
            style_table(table)
            for r_idx, row in enumerate(rows):
                for c_idx in range(cols):
                    cell = table.rows[r_idx].cells[c_idx]
                    cell.text = ""
                    p = cell.paragraphs[0]
                    val = row[c_idx] if c_idx < len(row) else ""
                    add_formatted_runs(p, val, size=9)
                    if r_idx == 0:
                        for run in p.runs:
                            run.bold = True
                            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            doc.add_paragraph()
            continue

        ul = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if ul:
            text = ul.group(2)
            p = doc.add_paragraph(style="List Bullet")
            add_formatted_runs(p, text)
            i += 1
            continue

        ol = re.match(r"^(\s*)\d+\.\s+(.*)$", line)
        if ol:
            text = ol.group(2)
            p = doc.add_paragraph(style="List Number")
            add_formatted_runs(p, text)
            i += 1
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        add_formatted_runs(p, line.strip())
        i += 1

    doc.save(str(docx_path))


def main():
    docs = Path(__file__).resolve().parent
    files = [
        "南沙电装OMS-开发文档.md",
        "南沙电装OMS-操作手册.md",
    ]
    for name in files:
        md = docs / name
        docx = docs / (md.stem + ".docx")
        convert(md, docx)
        print("wrote", docx)


if __name__ == "__main__":
    main()
